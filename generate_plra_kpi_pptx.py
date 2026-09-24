#!/usr/bin/env python3
"""
Generate a formal 4-slide PLRA KPI Project presentation.
Cover → Static/Dynamic Data → Figma & Documentation Status → Thank You
"""

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


# PLRA official green palette
GREEN = RGBColor(0x00, 0x68, 0x37)
GREEN_DARK = RGBColor(0x0B, 0x4F, 0x2C)
GREEN_MID = RGBColor(0x1B, 0x7A, 0x4A)
GREEN_SOFT = RGBColor(0xE8, 0xF5, 0xEE)
GREEN_BORDER = RGBColor(0x2F, 0x9E, 0x6B)
GOLD = RGBColor(0xC4, 0xA3, 0x5A)
GRAY = RGBColor(0x4A, 0x4A, 0x4A)
GRAY_LIGHT = RGBColor(0x88, 0x88, 0x88)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x22, 0x22, 0x22)


def set_run(run, size=14, bold=False, color=BLACK, italic=False):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def add_text(shape, lines, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}.get(valign, "t"),
        )
    except Exception:
        pass
    for i, item in enumerate(lines):
        if isinstance(item, str):
            text, size, bold, color, italic = item, 14, False, BLACK, False
        else:
            text, size, bold, color = item[0], item[1], item[2], item[3]
            italic = item[4] if len(item) > 4 else False
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = text
        set_run(run, size=size, bold=bold, color=color, italic=italic)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def fill_rect(slide, left, top, width, height, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    return sh


def rounded(slide, left, top, width, height, fill=WHITE, line=None, line_width=1.25, radius=0.08):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_width)
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    return sh


def oval(slide, left, top, size, fill=GREEN):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def plra_logo_mark(slide, left, top, size=0.85, dark=False):
    """Stylized PLRA emblem mark (green circle + P) for official branding."""
    fill = WHITE if dark else GREEN
    text_color = GREEN if dark else WHITE
    mark = oval(slide, Inches(left), Inches(top), Inches(size), fill=fill)
    add_text(mark, [("P", int(18 * size / 0.85), True, text_color)], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    return mark


def accent_line(slide, left, top, width=2.2):
    g = fill_rect(slide, Inches(left), Inches(top), Inches(width * 0.55), Inches(0.04), GOLD)
    gr = fill_rect(slide, Inches(left + width * 0.55), Inches(top), Inches(width * 0.45), Inches(0.04), GREEN)
    return g, gr


def set_fade_transition(slide):
    sld = slide._element
    for child in list(sld):
        if etree.QName(child).localname == "transition":
            sld.remove(child)
    tr = etree.Element(qn("p:transition"))
    tr.set("spd", "med")
    tr.set("advClick", "1")
    etree.SubElement(tr, qn("p:fade"))
    insert_at = 0
    for idx, child in enumerate(list(sld)):
        if etree.QName(child).localname in ("cSld", "clrMapOvr"):
            insert_at = idx + 1
    sld.insert(insert_at, tr)


def footer(slide, prs, light=False):
    color = RGBColor(0xB8, 0xD8, 0xC6) if light else GRAY_LIGHT
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.1), Inches(12.3), Inches(0.28))
    add_text(
        box,
        [("Punjab Land Records Authority (PLRA)  |  KPI Project", 10, False, color)],
        align=PP_ALIGN.CENTER,
    )


# -------------------- Slide 1: Cover --------------------

def slide_cover(prs):
    s = blank(prs)
    # Full green hero plane
    fill_rect(s, 0, 0, prs.slide_width, prs.slide_height, GREEN)
    # Soft accent band at bottom
    fill_rect(s, 0, Inches(6.35), prs.slide_width, Inches(1.15), GREEN_DARK)

    plra_logo_mark(s, 6.2, 1.05, size=1.05, dark=True)

    org = s.shapes.add_textbox(Inches(0.8), Inches(2.25), Inches(11.7), Inches(0.4))
    add_text(
        org,
        [("PUNJAB LAND RECORDS AUTHORITY", 14, True, RGBColor(0xD8, 0xF0, 0xE4))],
        align=PP_ALIGN.CENTER,
    )

    title = s.shapes.add_textbox(Inches(0.8), Inches(2.85), Inches(11.7), Inches(0.85))
    add_text(
        title,
        [("KPI", 54, True, WHITE)],
        align=PP_ALIGN.CENTER,
    )

    subtitle = s.shapes.add_textbox(Inches(0.8), Inches(3.65), Inches(11.7), Inches(0.55))
    add_text(
        subtitle,
        [("Key Performance Indicator Project", 26, True, WHITE)],
        align=PP_ALIGN.CENTER,
    )

    # Gold + white accent under title
    fill_rect(s, Inches(5.55), Inches(4.35), Inches(1.2), Inches(0.045), GOLD)
    fill_rect(s, Inches(6.75), Inches(4.35), Inches(1.0), Inches(0.045), WHITE)

    tag = s.shapes.add_textbox(Inches(1.5), Inches(4.6), Inches(10.3), Inches(0.7))
    add_text(
        tag,
        [
            ("District-wise Ranking of Authorities by Rank", 16, False, RGBColor(0xE0, 0xF2, 0xE8)),
            ("Performance Monitoring & Evaluation Platform", 14, False, RGBColor(0xB8, 0xD8, 0xC6)),
        ],
        align=PP_ALIGN.CENTER,
    )

    bottom = s.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(11.7), Inches(0.55))
    add_text(
        bottom,
        [
            ("Government of the Punjab", 13, True, WHITE),
            ("Status Briefing Presentation", 11, False, RGBColor(0xB8, 0xD8, 0xC6)),
        ],
        align=PP_ALIGN.CENTER,
    )
    return s


# -------------------- Slide 2: Static vs Dynamic Data --------------------

def data_column(slide, left, title, subtitle, items, accent=GREEN):
    width = 5.85
    height = 5.0
    rounded(slide, Inches(left), Inches(1.45), Inches(width), Inches(height), fill=WHITE, line=GREEN_BORDER, radius=0.06)
    fill_rect(slide, Inches(left), Inches(1.45), Inches(width), Inches(0.7), accent)

    ht = slide.shapes.add_textbox(Inches(left + 0.25), Inches(1.55), Inches(width - 0.5), Inches(0.5))
    add_text(ht, [(title, 18, True, WHITE)], align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE)

    st = slide.shapes.add_textbox(Inches(left + 0.3), Inches(2.3), Inches(width - 0.6), Inches(0.4))
    add_text(st, [(subtitle, 12, False, GRAY)])

    lines = [(f"•  {item}", 13, False, BLACK) for item in items]
    body = slide.shapes.add_textbox(Inches(left + 0.3), Inches(2.8), Inches(width - 0.6), Inches(3.4))
    add_text(body, lines)


def slide_data(prs):
    s = blank(prs)
    fill_rect(s, 0, 0, prs.slide_width, prs.slide_height, RGBColor(0xFC, 0xFC, 0xFC))
    fill_rect(s, 0, 0, prs.slide_width, Inches(0.12), GREEN)

    plra_logo_mark(s, 0.45, 0.32, size=0.42)
    brand = s.shapes.add_textbox(Inches(1.0), Inches(0.38), Inches(4), Inches(0.32))
    add_text(brand, [("PLRA  ·  KPI Project", 12, True, GREEN)])

    num = s.shapes.add_textbox(Inches(0.45), Inches(0.85), Inches(0.7), Inches(0.4))
    add_text(num, [("02", 22, True, GREEN)])
    title = s.shapes.add_textbox(Inches(1.15), Inches(0.85), Inches(11), Inches(0.4))
    add_text(title, [("Data Inputs — Static & Dynamic", 24, True, GREEN)])
    sub = s.shapes.add_textbox(Inches(1.15), Inches(1.2), Inches(11), Inches(0.25))
    add_text(sub, [("Reference data received statically vs. operational data received dynamically", 12, False, GRAY_LIGHT)])
    accent_line(s, 1.15, 1.48, width=1.8)

    data_column(
        s,
        0.55,
        "Static Data",
        "One-time / reference datasets received for setup",
        [
            "District & tehsil master lists",
            "Rank / designation hierarchy of authorities",
            "Officer-to-territory organizational mapping",
            "KPI indicator definitions & weightages",
            "Target / benchmark thresholds per rank",
            "Scoring & ranking formula configuration",
        ],
        accent=GREEN,
    )

    data_column(
        s,
        6.9,
        "Dynamic Data",
        "Ongoing / live feeds for performance measurement",
        [
            "Service delivery volumes (district-wise)",
            "Processing & turnaround time metrics",
            "Pendency and escalation statistics",
            "Complaint / grievance resolution indicators",
            "Period-wise performance score calculations",
            "Live district-wise authority rankings",
        ],
        accent=GREEN_DARK,
    )

    footer(s, prs)
    return s


# -------------------- Slide 3: Figma + Documentation --------------------

def status_card(slide, left, top, width, height, heading, lines):
    rounded(slide, Inches(left), Inches(top), Inches(width), Inches(height), fill=WHITE, line=GREEN_BORDER, radius=0.06)
    fill_rect(slide, Inches(left), Inches(top), Inches(0.1), Inches(height), GREEN)
    h = slide.shapes.add_textbox(Inches(left + 0.3), Inches(top + 0.2), Inches(width - 0.5), Inches(0.4))
    add_text(h, [(heading, 16, True, GREEN)])
    body = slide.shapes.add_textbox(Inches(left + 0.3), Inches(top + 0.7), Inches(width - 0.5), Inches(height - 0.9))
    add_text(body, [(ln, 13, False, BLACK) if not isinstance(ln, tuple) else ln for ln in lines])


def slide_status(prs):
    s = blank(prs)
    fill_rect(s, 0, 0, prs.slide_width, prs.slide_height, RGBColor(0xFC, 0xFC, 0xFC))
    fill_rect(s, 0, 0, prs.slide_width, Inches(0.12), GREEN)

    plra_logo_mark(s, 0.45, 0.32, size=0.42)
    brand = s.shapes.add_textbox(Inches(1.0), Inches(0.38), Inches(4), Inches(0.32))
    add_text(brand, [("PLRA  ·  KPI Project", 12, True, GREEN)])

    num = s.shapes.add_textbox(Inches(0.45), Inches(0.85), Inches(0.7), Inches(0.4))
    add_text(num, [("03", 22, True, GREEN)])
    title = s.shapes.add_textbox(Inches(1.15), Inches(0.85), Inches(11), Inches(0.4))
    add_text(title, [("Design Approval & Documentation Handover", 24, True, GREEN)])
    sub = s.shapes.add_textbox(Inches(1.15), Inches(1.2), Inches(11), Inches(0.25))
    add_text(sub, [("Current status communicated to the vendor", 12, False, GRAY_LIGHT)])
    accent_line(s, 1.15, 1.48, width=1.8)

    # Figma approved banner
    banner = rounded(s, Inches(0.55), Inches(1.75), Inches(12.2), Inches(1.35), fill=GREEN_SOFT, line=GREEN_BORDER, radius=0.06)
    fill_rect(s, Inches(0.55), Inches(1.75), Inches(0.12), Inches(1.35), GREEN)
    badge = rounded(s, Inches(0.9), Inches(2.05), Inches(1.6), Inches(0.45), fill=GREEN, radius=0.3)
    add_text(badge, [("APPROVED", 12, True, WHITE)], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    ft = s.shapes.add_textbox(Inches(2.7), Inches(1.95), Inches(9.5), Inches(0.4))
    add_text(ft, [("Final Figma Design — Approved", 18, True, GREEN)])
    fd = s.shapes.add_textbox(Inches(2.7), Inches(2.4), Inches(9.5), Inches(0.45))
    add_text(
        fd,
        [("UI/UX design for the KPI district-wise ranking platform has been finalized and formally approved.", 13, False, GRAY)],
    )

    # Documentation email card
    status_card(
        s,
        0.55,
        3.35,
        12.2,
        3.35,
        "Documentation & Project Plan — Shared with Vendor",
        [
            ("The PLRA KPI team has sent an email to the vendor regarding the following deliverables:", 13, False, GRAY),
            ("", 6, False, WHITE),
            ("•  SRS  — Software Requirements Specification", 14, False, BLACK),
            ("•  BRD  — Business Requirements Document", 14, False, BLACK),
            ("•  Workflow  — End-to-end process flows", 14, False, BLACK),
            ("•  User Manual  — End-user operational guide", 14, False, BLACK),
            ("•  Project Plan  — Timeline, milestones & delivery schedule", 14, False, BLACK),
        ],
    )

    footer(s, prs)
    return s


# -------------------- Slide 4: Thank You --------------------

def slide_thanks(prs):
    s = blank(prs)
    fill_rect(s, 0, 0, prs.slide_width, prs.slide_height, WHITE)
    fill_rect(s, 0, Inches(2.0), prs.slide_width, Inches(3.5), GREEN)

    plra_logo_mark(s, 6.2, 2.25, size=0.75, dark=True)

    t = s.shapes.add_textbox(Inches(0.8), Inches(3.2), Inches(11.7), Inches(0.7))
    add_text(t, [("Thank You", 44, True, WHITE)], align=PP_ALIGN.CENTER)

    fill_rect(s, Inches(5.9), Inches(3.95), Inches(0.8), Inches(0.04), GOLD)
    fill_rect(s, Inches(6.7), Inches(3.95), Inches(0.7), Inches(0.04), WHITE)

    d = s.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(11.7), Inches(0.4))
    add_text(
        d,
        [("KPI — Key Performance Indicator Project", 16, False, RGBColor(0xD8, 0xF0, 0xE4))],
        align=PP_ALIGN.CENTER,
    )

    f = s.shapes.add_textbox(Inches(0.8), Inches(4.7), Inches(11.7), Inches(0.35))
    add_text(
        f,
        [("Punjab Land Records Authority (PLRA)  |  Government of the Punjab", 13, True, WHITE)],
        align=PP_ALIGN.CENTER,
    )

    footer(s, prs, light=False)
    return s


def validate_pptx(path):
    import zipfile
    from pptx import Presentation as P

    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
        if bad:
            raise RuntimeError(f"Corrupt zip entry: {bad}")
        for name in z.namelist():
            if name.endswith(".xml") or name.endswith(".rels"):
                etree.fromstring(z.read(name))
    return len(P(path).slides)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide_cover(prs)
    slide_data(prs)
    slide_status(prs)
    slide_thanks(prs)

    for slide in prs.slides:
        set_fade_transition(slide)

    out = Path("/workspace/PLRA_KPI_Project_Presentation.pptx")
    tmp = Path("/workspace/PLRA_KPI_Project_Presentation.tmp.pptx")
    prs.save(str(tmp))
    n = validate_pptx(str(tmp))
    tmp.replace(out)
    print(f"Saved {out} with {n} slides (validated)")
    return str(out)


if __name__ == "__main__":
    build()
