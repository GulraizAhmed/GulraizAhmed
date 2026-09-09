#!/usr/bin/env python3
"""
Generate a valid formal PLRA LMS PowerPoint presentation.
Uses only well-formed OOXML transitions/animations and standard shapes.
"""

from copy import deepcopy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


# Brand colors matching reference slides
GREEN = RGBColor(0x00, 0x68, 0x37)
GREEN_DARK = RGBColor(0x0B, 0x4F, 0x2C)
GREEN_MID = RGBColor(0x1B, 0x7A, 0x4A)
GREEN_SOFT = RGBColor(0xE8, 0xF5, 0xEE)
GREEN_BORDER = RGBColor(0x2F, 0x9E, 0x6B)
GRAY = RGBColor(0x4A, 0x4A, 0x4A)
GRAY_LIGHT = RGBColor(0x88, 0x88, 0x88)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x22, 0x22, 0x22)
BG = RGBColor(0xFC, 0xFC, 0xFC)

NSMAP = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


def set_run(run, size=14, bold=False, color=BLACK, italic=False):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def add_text(shape, lines, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    try:
        anchor = {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}.get(valign, "t")
        tf._txBody.bodyPr.set("anchor", anchor)
    except Exception:
        pass
    for i, item in enumerate(lines):
        if isinstance(item, str):
            text, size, bold, color, italic = item, 14, False, BLACK, False
        else:
            text, size, bold, color = item[0], item[1], item[2], item[3]
            italic = item[4] if len(item) > 4 else False
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = text
        set_run(run, size=size, bold=bold, color=color, italic=italic)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def fill_bg(slide, prs, color=BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def rounded(slide, left, top, width, height, fill=WHITE, line=None, line_width=1.25, radius=0.08):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_width)
    try:
        shape.adjustments[0] = radius
    except Exception:
        pass
    return shape


def oval(slide, left, top, size, fill=GREEN):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def dots(slide, left, top, cols=7, rows=2, size=0.055, gap=0.13):
    for r in range(rows):
        for c in range(cols):
            s = oval(
                slide,
                Inches(left + c * gap),
                Inches(top + r * gap),
                Inches(size),
                fill=GREEN_BORDER,
            )
            # soft look
            s.fill.solid()
            s.fill.fore_color.rgb = RGBColor(0xA8, 0xD5, 0xB5)


def brand_header(slide, text="PLRA - LMS"):
    logo = oval(slide, Inches(0.45), Inches(0.28), Inches(0.4), GREEN)
    add_text(logo, [("P", 12, True, WHITE)], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    box = slide.shapes.add_textbox(Inches(0.95), Inches(0.32), Inches(4.5), Inches(0.35))
    add_text(box, [(text, 15, True, GREEN)])


def title_text(slide, title, y=0.85):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(y), Inches(12.2), Inches(0.55))
    add_text(box, [(title.upper(), 28, True, GREEN)])


def corner_accents(slide, prs):
    # Top-right accent
    tr = rounded(
        slide,
        prs.slide_width - Inches(2.2),
        Inches(-0.55),
        Inches(2.8),
        Inches(1.55),
        fill=GREEN,
        radius=0.22,
    )
    # Bottom-left bars
    bar1 = slide.shapes.add_shape(
        MSO_SHAPE.PARALLELOGRAM,
        Inches(-0.7),
        prs.slide_height - Inches(1.0),
        Inches(7.2),
        Inches(0.48),
    )
    bar1.fill.solid()
    bar1.fill.fore_color.rgb = GREEN
    bar1.line.fill.background()

    bar2 = slide.shapes.add_shape(
        MSO_SHAPE.PARALLELOGRAM,
        Inches(1.0),
        prs.slide_height - Inches(0.65),
        Inches(7.8),
        Inches(0.4),
    )
    bar2.fill.solid()
    bar2.fill.fore_color.rgb = GREEN_MID
    bar2.line.fill.background()

    # Bottom-right accent
    rounded(
        slide,
        prs.slide_width - Inches(1.45),
        prs.slide_height - Inches(1.2),
        Inches(2.0),
        Inches(1.6),
        fill=GREEN_DARK,
        radius=0.2,
    )


def accent_panel(slide, left, top, width, height, label="LMS"):
    """Decorative green panel instead of potentially problematic external images."""
    frame = rounded(slide, Inches(left), Inches(top), Inches(width), Inches(height), fill=GREEN, radius=0.1)
    inner = rounded(
        slide,
        Inches(left + 0.14),
        Inches(top + 0.14),
        Inches(width - 0.28),
        Inches(height - 0.28),
        fill=GREEN_SOFT,
        radius=0.1,
    )
    # content blocks to suggest a dashboard
    for i, (w, h, dx, dy) in enumerate(
        [
            (1.35, 0.7, 0.35, 0.45),
            (1.35, 0.7, 1.9, 0.45),
            (2.9, 1.5, 0.35, 1.4),
            (2.9, 1.1, 0.35, 3.1),
        ]
    ):
        card = rounded(
            slide,
            Inches(left + dx),
            Inches(top + dy),
            Inches(w),
            Inches(h),
            fill=WHITE,
            line=GREEN_BORDER,
            radius=0.08,
        )
    lbl = slide.shapes.add_textbox(Inches(left + 0.3), Inches(top + height - 0.7), Inches(width - 0.6), Inches(0.4))
    add_text(lbl, [(label, 14, True, GREEN)], align=PP_ALIGN.CENTER)
    return frame


# Unique animation timing node IDs across a presentation build
_ANIM_ID = 10


def _next_anim_ids(count=3):
    global _ANIM_ID
    ids = list(range(_ANIM_ID, _ANIM_ID + count))
    _ANIM_ID += count
    return ids


def _normalize_slide_order(sld):
    """Ensure OOXML child order: cSld, clrMapOvr, transition, timing, ..."""
    children = list(sld)
    buckets = {"cSld": None, "clrMapOvr": None, "transition": None, "timing": None}
    others = []
    for child in children:
        local = etree.QName(child).localname
        if local in buckets and buckets[local] is None:
            buckets[local] = child
        else:
            others.append(child)
    for child in children:
        sld.remove(child)
    for key in ("cSld", "clrMapOvr", "transition", "timing"):
        if buckets[key] is not None:
            sld.append(buckets[key])
    for child in others:
        sld.append(child)


def set_transition(slide, kind="fade", direction="l"):
    """Insert/replace a valid slide transition; keep existing timing."""
    sld = slide._element
    for child in list(sld):
        if etree.QName(child).localname == "transition":
            sld.remove(child)

    transition = etree.Element(qn("p:transition"))
    transition.set("spd", "med")
    transition.set("advClick", "1")

    dirs = {"l", "r", "u", "d"}
    d = direction if direction in dirs else "l"

    if kind == "push":
        node = etree.SubElement(transition, qn("p:push"))
        node.set("dir", d)
    elif kind == "wipe":
        node = etree.SubElement(transition, qn("p:wipe"))
        node.set("dir", d)
    elif kind == "cover":
        node = etree.SubElement(transition, qn("p:cover"))
        node.set("dir", d)
    else:
        etree.SubElement(transition, qn("p:fade"))

    sld.append(transition)
    _normalize_slide_order(sld)


def add_fade_animation(slide, shape, order=None):
    """
    Add a Fade entrance animation on click for a shape.
    Builds a minimal valid p:timing tree if missing.
    Uses globally unique timing node IDs (required by PowerPoint).
    """
    global _ANIM_ID
    sld = slide._element
    timing = None
    for child in sld:
        if etree.QName(child).localname == "timing":
            timing = child
            break

    if timing is None:
        timing = etree.Element(qn("p:timing"))
        sld.append(timing)

        tn_lst = etree.SubElement(timing, qn("p:tnLst"))
        par = etree.SubElement(tn_lst, qn("p:par"))
        ctn = etree.SubElement(par, qn("p:cTn"))
        ctn.set("id", "1")
        ctn.set("dur", "indefinite")
        ctn.set("restart", "never")
        ctn.set("nodeType", "tmRoot")
        child_tn = etree.SubElement(ctn, qn("p:childTnLst"))
        seq = etree.SubElement(child_tn, qn("p:seq"))
        seq.set("concurrent", "1")
        seq.set("nextAc", "seek")
        seq_ctn = etree.SubElement(seq, qn("p:cTn"))
        seq_ctn.set("id", "2")
        seq_ctn.set("dur", "indefinite")
        seq_ctn.set("nodeType", "mainSeq")
        etree.SubElement(seq_ctn, qn("p:childTnLst"))
        prev = etree.SubElement(seq, qn("p:prevCondLst"))
        cond = etree.SubElement(prev, qn("p:cond"))
        cond.set("evt", "onPrev")
        cond.set("delay", "0")
        tgt = etree.SubElement(cond, qn("p:tgtEl"))
        etree.SubElement(tgt, qn("p:sldTgt"))
        nxt = etree.SubElement(seq, qn("p:nextCondLst"))
        cond2 = etree.SubElement(nxt, qn("p:cond"))
        cond2.set("evt", "onNext")
        cond2.set("delay", "0")
        tgt2 = etree.SubElement(cond2, qn("p:tgtEl"))
        etree.SubElement(tgt2, qn("p:sldTgt"))
        _ANIM_ID = max(_ANIM_ID, 10)

    # Find mainSeq childTnLst
    p_ns = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    main_list = None
    for ctn in timing.findall(f".//{p_ns}cTn"):
        if ctn.get("nodeType") == "mainSeq":
            main_list = ctn.find(f"{p_ns}childTnLst")
            break
    if main_list is None:
        return

    id1, id2, id3 = _next_anim_ids(3)
    spid = str(shape.shape_id)

    par = etree.SubElement(main_list, qn("p:par"))
    ctn = etree.SubElement(par, qn("p:cTn"))
    ctn.set("id", str(id1))
    ctn.set("fill", "hold")
    ctn.set("presetID", "10")
    ctn.set("presetClass", "entr")
    ctn.set("presetSubtype", "0")
    ctn.set("grpId", "0")
    ctn.set("nodeType", "clickEffect")
    st = etree.SubElement(ctn, qn("p:stCondLst"))
    cond = etree.SubElement(st, qn("p:cond"))
    cond.set("delay", "0")
    child = etree.SubElement(ctn, qn("p:childTnLst"))

    # set visibility
    set_el = etree.SubElement(child, qn("p:set"))
    cBhvr = etree.SubElement(set_el, qn("p:cBhvr"))
    cTn2 = etree.SubElement(cBhvr, qn("p:cTn"))
    cTn2.set("id", str(id2))
    cTn2.set("dur", "1")
    cTn2.set("fill", "hold")
    st2 = etree.SubElement(cTn2, qn("p:stCondLst"))
    cond2 = etree.SubElement(st2, qn("p:cond"))
    cond2.set("delay", "0")
    tgt = etree.SubElement(cBhvr, qn("p:tgtEl"))
    sp = etree.SubElement(tgt, qn("p:spTgt"))
    sp.set("spid", spid)
    attr = etree.SubElement(cBhvr, qn("p:attrNameLst"))
    an = etree.SubElement(attr, qn("p:attrName"))
    an.text = "style.visibility"
    to = etree.SubElement(set_el, qn("p:to"))
    val = etree.SubElement(to, qn("p:strVal"))
    val.set("val", "visible")

    # fade effect
    anim = etree.SubElement(child, qn("p:animEffect"))
    anim.set("transition", "in")
    anim.set("filter", "fade")
    cBhvr2 = etree.SubElement(anim, qn("p:cBhvr"))
    cTn3 = etree.SubElement(cBhvr2, qn("p:cTn"))
    cTn3.set("id", str(id3))
    cTn3.set("dur", "500")
    tgt2 = etree.SubElement(cBhvr2, qn("p:tgtEl"))
    sp2 = etree.SubElement(tgt2, qn("p:spTgt"))
    sp2.set("spid", spid)

    _normalize_slide_order(sld)


def arrow_bullets(slide, items, left=0.55, top=1.65, width=7.4, size=16, animate=False):
    shapes = []
    y = top
    for idx, it in enumerate(items, 1):
        circ = oval(slide, Inches(left), Inches(y), Inches(0.3), GREEN)
        add_text(circ, [(">", 12, True, WHITE)], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        tb = slide.shapes.add_textbox(Inches(left + 0.45), Inches(y - 0.02), Inches(width), Inches(0.5))
        add_text(tb, [(it, size, False, GRAY)])
        if animate:
            add_fade_animation(slide, tb)
        shapes.append(circ)
        y += 0.7
    return shapes


def card_grid(slide, cards, animate=True):
    """cards: list of (title, bullets) in a row."""
    n = len(cards)
    gap = 0.2
    total_w = 12.4
    card_w = (total_w - gap * (n - 1)) / n
    shapes = []
    for i, (title, bullets) in enumerate(cards):
        x = 0.45 + i * (card_w + gap)
        card = rounded(
            slide,
            Inches(x),
            Inches(1.55),
            Inches(card_w),
            Inches(4.55),
            fill=WHITE,
            line=GREEN_BORDER,
            line_width=1.5,
        )
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(1.55), Inches(0.1), Inches(4.55))
        bar.fill.solid()
        bar.fill.fore_color.rgb = GREEN
        bar.line.fill.background()
        ht = slide.shapes.add_textbox(Inches(x + 0.25), Inches(1.75), Inches(card_w - 0.4), Inches(0.45))
        add_text(ht, [(title, 14, True, GREEN)])
        bt = slide.shapes.add_textbox(Inches(x + 0.25), Inches(2.35), Inches(card_w - 0.4), Inches(3.5))
        add_text(bt, [(f"•  {b}", 12, False, GRAY) for b in bullets])
        if animate:
            add_fade_animation(slide, card)
        shapes.append(card)
    return shapes


# -------------------- slides --------------------

def slide_title(prs):
    s = blank(prs)
    fill_bg(s, prs, WHITE)
    corner_accents(s, prs)
    dots(s, 5.7, 0.35)

    logo = oval(s, Inches(0.55), Inches(0.4), Inches(0.48), GREEN)
    add_text(logo, [("P", 14, True, WHITE)], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    box = s.shapes.add_textbox(Inches(1.15), Inches(0.48), Inches(3), Inches(0.35))
    add_text(box, [("PLRA", 18, True, GREEN)])

    box = s.shapes.add_textbox(Inches(0.55), Inches(2.05), Inches(7.8), Inches(1.0))
    add_text(box, [("LMS", 64, True, GREEN)])
    add_fade_animation(s, box)
    box = s.shapes.add_textbox(Inches(0.55), Inches(3.15), Inches(7.8), Inches(0.55))
    add_text(box, [("Learning Management System", 24, True, GREEN)])
    add_fade_animation(s, box)

    banner = rounded(s, Inches(0.55), Inches(5.35), Inches(6.8), Inches(0.65), fill=GREEN, radius=0.4)
    add_text(
        banner,
        [("Website  |  lms.punjab-zameen.gov.pk", 15, True, WHITE)],
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_fade_animation(s, banner)
    accent_panel(s, 8.55, 1.55, 4.2, 4.4, "PULSE LMS")
    return s


def slide_agenda(prs):
    s = blank(prs)
    fill_bg(s, prs)
    corner_accents(s, prs)
    brand_header(s)
    dots(s, 6.0, 0.28)
    title_text(s, "Agenda")

    left = [
        "01  Project Background",
        "02  Key Objectives",
        "03  Purpose & Scope",
        "04  Types of Users",
        "05  Role Dashboards",
        "06  Core Platform Capabilities",
    ]
    right = [
        "07  Learning & Assessment",
        "08  Certification & Analytics",
        "09  Application Workflows",
        "10  Technology Stack",
        "11  Operating Environment",
        "12  Closing",
    ]
    cards = []
    for i, items in enumerate([left, right]):
        x = 0.55 if i == 0 else 6.95
        card = rounded(s, Inches(x), Inches(1.55), Inches(5.85), Inches(4.55), fill=WHITE, line=GREEN_BORDER)
        tb = s.shapes.add_textbox(Inches(x + 0.35), Inches(1.8), Inches(5.2), Inches(4.1))
        add_text(tb, [(it, 15, False, GRAY) for it in items])
        cards.append(card)
        add_fade_animation(s, card, i + 1)
    return s


def slide_background(prs):
    s = blank(prs)
    fill_bg(s, prs)
    corner_accents(s, prs)
    brand_header(s)
    dots(s, 6.0, 0.28)
    title_text(s, "Project Background")

    paras = [
        "LMS was developed to modernize and standardize learning across PLRA by providing a centralized digital platform for training, assessment, certification, and compliance management.",
        "Traditional training processes often rely on manual coordination, fragmented records, and limited visibility into learner progress. LMS addresses these challenges by delivering a secure, scalable, and data-driven learning ecosystem that supports administrators, trainers, and learners through a single platform.",
    ]
    tb = s.shapes.add_textbox(Inches(0.55), Inches(1.55), Inches(7.5), Inches(4.4))
    lines = []
    for p in paras:
        lines.append((p, 15, False, GRAY))
        lines.append(("", 10, False, GRAY))
    add_text(tb, lines)
    add_fade_animation(s, tb)
    accent_panel(s, 8.4, 1.45, 4.35, 4.5, "Centralized Learning")
    return s


def slide_objectives(prs):
    s = blank(prs)
    fill_bg(s, prs)
    corner_accents(s, prs)
    brand_header(s)
    dots(s, 6.0, 0.28)
    title_text(s, "Key Objectives")
    items = [
        "Standardize training delivery across all departments and regions",
        "Digitize assessments, certifications, and learner records",
        "Improve visibility into training progress and compliance status",
        "Reduce administrative effort through automation and self-service workflows",
    ]
    arrow_bullets(s, items, left=0.55, top=1.7, width=7.5, size=16, animate=True)
    accent_panel(s, 8.5, 1.4, 4.25, 4.45, "Objectives")
    return s


def slide_purpose_scope(prs):
    s = blank(prs)
    fill_bg(s, prs)
    corner_accents(s, prs)
    brand_header(s)
    dots(s, 6.0, 0.28)
    title_text(s, "Purpose & Scope")

    c1 = rounded(s, Inches(0.5), Inches(1.45), Inches(12.3), Inches(1.65), fill=GREEN_SOFT, line=GREEN_BORDER)
    t = s.shapes.add_textbox(Inches(0.75), Inches(1.55), Inches(11.8), Inches(0.3))
    add_text(t, [("PURPOSE", 13, True, GREEN)])
    b = s.shapes.add_textbox(Inches(0.75), Inches(1.95), Inches(11.8), Inches(0.95))
    add_text(
        b,
        [
            (
                "PULSE LMS is a Laravel-based Learning Management System for PLRA with CIMS authentication, "
                "role-based access, trainee onboarding, course delivery, quizzes, certificates, reports, "
                "notifications, support tickets, chat, forums, and protected media access.",
                13,
                False,
                GRAY,
            )
        ],
    )
    add_fade_animation(s, c1, 1)

    blocks = [
        (
            "IN SCOPE",
            [
                "CIMS login, RBAC, onboarding & approvals",
                "Courses, modules, lessons, quizzes & media",
                "Progress, certificates, badges & reports",
                "Alerts, support tickets, chat & forums",
            ],
        ),
        (
            "OUT OF SCOPE",
            [
                "Authoring of legal training content itself",
                "Video / translation / trainer-guide production",
                "Content creation as a separate project deliverable",
                "LMS stores and delivers materials once ready",
            ],
        ),
    ]
    for i, (title, bullets) in enumerate(blocks):
        x = 0.5 + i * 6.35
        card = rounded(s, Inches(x), Inches(3.35), Inches(6.1), Inches(2.85), fill=WHITE, line=GREEN_BORDER)
        ht = s.shapes.add_textbox(Inches(x + 0.25), Inches(3.5), Inches(5.5), Inches(0.35))
        add_text(ht, [(title, 14, True, GREEN)])
        bt = s.shapes.add_textbox(Inches(x + 0.25), Inches(4.0), Inches(5.5), Inches(1.9))
        add_text(bt, [(f"•  {b}", 13, False, GRAY) for b in bullets])
        add_fade_animation(s, card, i + 2)
    return s


def slide_users(prs):
    s = blank(prs)
    fill_bg(s, prs)
    corner_accents(s, prs)
    brand_header(s)
    dots(s, 6.0, 0.28)
    title_text(s, "Types of Users")

    users = [
        ("01", "Super Admin", "Platform governance & control"),
        ("02", "Admin", "Learning operations & management"),
        ("03", "Trainer", "Course delivery & learner insights"),
        ("04", "Trainee", "Personalized learning journey"),
    ]
    for i, (num, name, desc) in enumerate(users):
        x = 0.55 + i * 3.15
        card = rounded(s, Inches(x), Inches(2.0), Inches(2.95), Inches(3.45), fill=WHITE, line=GREEN_BORDER, line_width=1.75)
        circ = oval(s, Inches(x + 0.95), Inches(2.35), Inches(1.0), GREEN)
        add_text(circ, [(num, 18, True, WHITE)], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        nt = s.shapes.add_textbox(Inches(x + 0.15), Inches(3.6), Inches(2.65), Inches(0.4))
        add_text(nt, [(name, 15, True, GREEN)], align=PP_ALIGN.CENTER)
        dt = s.shapes.add_textbox(Inches(x + 0.15), Inches(4.15), Inches(2.65), Inches(0.9))
        add_text(dt, [(desc, 12, False, GRAY)], align=PP_ALIGN.CENTER)
        add_fade_animation(s, card, i + 1)
    return s


def role_slide(prs, title, subtitle, intro, responsibilities, panel_label):
    s = blank(prs)
    fill_bg(s, prs)
    corner_accents(s, prs)
    brand_header(s)
    dots(s, 5.8, 0.28)

    t = s.shapes.add_textbox(Inches(0.55), Inches(0.85), Inches(7.8), Inches(0.45))
    add_text(t, [(title.upper(), 26, True, GREEN)])
    st = s.shapes.add_textbox(Inches(0.55), Inches(1.35), Inches(7.8), Inches(0.35))
    add_text(st, [(subtitle, 14, True, GREEN_MID)])
    intro_box = s.shapes.add_textbox(Inches(0.55), Inches(1.8), Inches(7.5), Inches(1.05))
    add_text(intro_box, [(intro, 13, False, GRAY)])

    kr = s.shapes.add_textbox(Inches(0.55), Inches(2.95), Inches(7.5), Inches(0.3))
    add_text(kr, [("Key Responsibilities", 13, True, GREEN)])

    y = 3.35
    for idx, item in enumerate(responsibilities, 1):
        circ = oval(s, Inches(0.55), Inches(y), Inches(0.24), GREEN)
        add_text(circ, [("+", 9, True, WHITE)], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        tb = s.shapes.add_textbox(Inches(0.95), Inches(y - 0.02), Inches(7.0), Inches(0.32))
        add_text(tb, [(item, 12, False, GRAY)])
        add_fade_animation(s, circ, idx)
        y += 0.4

    accent_panel(s, 8.45, 1.3, 4.3, 4.55, panel_label)
    return s


def workflow_slide(prs, title, workflows):
    s = blank(prs)
    fill_bg(s, prs)
    corner_accents(s, prs)
    brand_header(s)
    dots(s, 6.0, 0.28)
    title_text(s, title)

    positions = [(0.45, 1.45), (6.85, 1.45), (0.45, 4.1), (6.85, 4.1)]
    for i, (name, steps) in enumerate(workflows[:4]):
        x, y = positions[i]
        card = rounded(s, Inches(x), Inches(y), Inches(6.05), Inches(2.4), fill=WHITE, line=GREEN_BORDER)
        circ = oval(s, Inches(x + 0.2), Inches(y + 0.2), Inches(0.36), GREEN)
        add_text(circ, [(f"{i+1:02d}", 11, True, WHITE)], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        ht = s.shapes.add_textbox(Inches(x + 0.7), Inches(y + 0.22), Inches(5.1), Inches(0.35))
        add_text(ht, [(name, 14, True, GREEN)])
        bt = s.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.7), Inches(5.55), Inches(1.55))
        add_text(bt, [(f"•  {st}", 11, False, GRAY) for st in steps])
        add_fade_animation(s, card, i + 1)
    return s


def slide_support(prs):
    s = blank(prs)
    fill_bg(s, prs)
    corner_accents(s, prs)
    brand_header(s)
    dots(s, 6.0, 0.28)
    title_text(s, "Support & Communication Workflow")
    steps = [
        ("01", "Raise", "User creates a support ticket or starts chat / forum interaction."),
        ("02", "Track", "System stores messages, participants, attachments, and ticket status."),
        ("03", "Resolve", "Authorized users respond, update status, or moderate where permitted."),
    ]
    for i, (num, title, desc) in enumerate(steps):
        x = 0.7 + i * 4.15
        card = rounded(s, Inches(x), Inches(2.15), Inches(3.9), Inches(3.5), fill=WHITE, line=GREEN_BORDER)
        circ = oval(s, Inches(x + 1.4), Inches(2.5), Inches(1.05), GREEN)
        add_text(circ, [(num, 18, True, WHITE)], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        nt = s.shapes.add_textbox(Inches(x + 0.2), Inches(3.8), Inches(3.5), Inches(0.4))
        add_text(nt, [(title, 17, True, GREEN)], align=PP_ALIGN.CENTER)
        dt = s.shapes.add_textbox(Inches(x + 0.25), Inches(4.35), Inches(3.4), Inches(1.0))
        add_text(dt, [(desc, 12, False, GRAY)], align=PP_ALIGN.CENTER)
        add_fade_animation(s, card, i + 1)
    return s


def slide_tech(prs):
    s = blank(prs)
    fill_bg(s, prs)
    corner_accents(s, prs)
    brand_header(s)
    dots(s, 6.0, 0.28)
    title_text(s, "Technology Stack & Operating Environment")

    intro = s.shapes.add_textbox(Inches(0.55), Inches(1.4), Inches(12.2), Inches(0.7))
    add_text(
        intro,
        [
            (
                "LMS is a server-rendered Laravel web application organized by domains. It runs with PHP-FPM and Nginx, "
                "uses a relational database, integrates with CIMS for authentication, and stores media via S3-compatible storage (MinIO).",
                13,
                False,
                GRAY,
            )
        ],
    )

    stacks = [
        ("Application", ["Laravel (PHP 8.2+)", "Blade role-specific UI", "Composer dependencies", "Node / Vite assets"]),
        ("Runtime", ["Docker containers", "Nginx reverse proxy", "PHP-FPM app runtime", "Queue & scheduler workers"]),
        ("Data & Storage", ["Relational database", "S3 / MinIO object storage", "Persistent volumes", "CSV reporting data"]),
        ("Integrations", ["CIMS auth (PKCE)", "SMTP / mail services", "HTTPS termination", "Modern browser clients"]),
    ]
    for i, (title, bullets) in enumerate(stacks):
        x = 0.45 + i * 3.2
        card = rounded(s, Inches(x), Inches(2.3), Inches(3.05), Inches(3.7), fill=WHITE, line=GREEN_BORDER)
        head = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(2.3), Inches(3.05), Inches(0.5))
        head.fill.solid()
        head.fill.fore_color.rgb = GREEN
        head.line.fill.background()
        ht = s.shapes.add_textbox(Inches(x + 0.1), Inches(2.38), Inches(2.85), Inches(0.35))
        add_text(ht, [(title, 13, True, WHITE)], align=PP_ALIGN.CENTER)
        bt = s.shapes.add_textbox(Inches(x + 0.2), Inches(3.0), Inches(2.7), Inches(2.7))
        add_text(bt, [(f"•  {b}", 12, False, GRAY) for b in bullets])
        add_fade_animation(s, card, i + 1)
    return s


def slide_closing(prs):
    s = blank(prs)
    fill_bg(s, prs, WHITE)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.05), prs.slide_width, Inches(3.3))
    band.fill.solid()
    band.fill.fore_color.rgb = GREEN
    band.line.fill.background()
    dots(s, 5.5, 0.4)

    logo = oval(s, Inches(6.3), Inches(2.3), Inches(0.65), WHITE)
    add_text(logo, [("P", 16, True, GREEN)], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)

    t = s.shapes.add_textbox(Inches(0.8), Inches(3.15), Inches(11.7), Inches(0.65))
    add_text(t, [("Thank You", 40, True, WHITE)], align=PP_ALIGN.CENTER)
    add_fade_animation(s, t)
    d = s.shapes.add_textbox(Inches(0.8), Inches(3.9), Inches(11.7), Inches(0.4))
    add_text(
        d,
        [("PLRA LMS — Modernizing Learning Across Punjab Land Records Authority", 14, False, RGBColor(0xD8, 0xF0, 0xE4))],
        align=PP_ALIGN.CENTER,
    )
    add_fade_animation(s, d)
    f = s.shapes.add_textbox(Inches(0.8), Inches(4.45), Inches(11.7), Inches(0.35))
    add_text(
        f,
        [("lms.punjab-zameen.gov.pk  |  Punjab Land Records Authority (PLRA)", 12, True, WHITE)],
        align=PP_ALIGN.CENTER,
    )
    add_fade_animation(s, f)
    return s


def validate_pptx(path):
    import zipfile
    from pptx import Presentation as P

    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
        if bad:
            raise RuntimeError(f"Corrupt zip entry: {bad}")
        slide_xmls = [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        if not slide_xmls:
            raise RuntimeError("No slides found in package")
        for name in z.namelist():
            if name.endswith(".xml") or name.endswith(".rels"):
                etree.fromstring(z.read(name))
        # Check slide XML has unique timing IDs and valid structure
        p_ns = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
        for name in sorted(slide_xmls):
            root = etree.fromstring(z.read(name))
            ids = [c.get("id") for c in root.findall(f".//{p_ns}cTn") if c.get("id")]
            if len(ids) != len(set(ids)):
                raise RuntimeError(f"Duplicate timing IDs in {name}: {ids}")
            locals_ = [etree.QName(c).localname for c in root]
            if "cSld" not in locals_:
                raise RuntimeError(f"Missing cSld in {name}")
            # transition should appear before timing when both exist
            if "transition" in locals_ and "timing" in locals_:
                if locals_.index("transition") > locals_.index("timing"):
                    raise RuntimeError(f"Invalid element order in {name}: timing before transition")
    prs = P(path)
    if len(prs.slides) < 1:
        raise RuntimeError("Presentation has no slides")
    return len(prs.slides)


def build():
    global _ANIM_ID
    _ANIM_ID = 10

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slides = []
    slides.append(slide_title(prs))
    slides.append(slide_agenda(prs))
    slides.append(slide_background(prs))
    slides.append(slide_objectives(prs))
    slides.append(slide_purpose_scope(prs))
    slides.append(slide_users(prs))

    slides.append(
        role_slide(
            prs,
            "Super Admin",
            "Complete Platform Governance & Control",
            "The Super Admin dashboard provides centralized oversight of the entire LMS ecosystem, enabling efficient management of administrators and trainers.",
            [
                "Manage Admin & Trainer Accounts",
                "Monitor Platform Health & System Resources",
                "Configure Roles, Permissions & Access Controls",
                "Oversee Alerts, Compliance & Governance",
            ],
            "Super Admin",
        )
    )
    slides.append(
        role_slide(
            prs,
            "Admin",
            "Centralized Learning Operations & Management",
            "The Admin dashboard serves as the operational hub of LMS for courses, trainers, trainees, approvals, and platform activities.",
            [
                "Create & Manage Courses, Categories & Topics",
                "Review & Approve Trainee Registrations",
                "Manage Trainers, Ranks & Learning Assignments",
                "Publish Alerts & System Notifications",
                "Oversee Media Library & Learning Resources",
                "Monitor Support Tickets & User Activity",
            ],
            "Admin Hub",
        )
    )
    slides.append(
        role_slide(
            prs,
            "Trainer",
            "Empowering Instructors Through Data & Insights",
            "The Trainer dashboard provides a workspace for managing assigned courses, monitoring learner progress, and tracking performance outcomes.",
            [
                "Manage Assigned Courses & Learning Activities",
                "Monitor Student Progress & Course Completion",
                "Analyze Learner Performance & Assessment Results",
                "Identify At-Risk Learners Through Analytics",
                "Respond to Course-Related Support Tickets",
                "Improve Outcomes Through Data-Driven Insights",
            ],
            "Trainer",
        )
    )
    slides.append(
        role_slide(
            prs,
            "Trainee",
            "Your Learning. Your Progress. Your Growth.",
            "The Trainee dashboard provides a personalized learning experience to access courses, track progress, manage activities, and achieve certification goals.",
            [
                "Access Enrolled Courses & Learning Materials",
                "Track Lesson Progress & Course Completion",
                "Monitor Study Hours & Learning Performance",
                "Stay Updated with Alerts & Notifications",
                "Manage Calendar, Deadlines & Learning Tasks",
                "Create & Track Support Tickets",
                "Earn Certificates Upon Course Completion",
            ],
            "Trainee",
        )
    )

    # capabilities
    s = blank(prs)
    fill_bg(s, prs)
    corner_accents(s, prs)
    brand_header(s)
    dots(s, 6.0, 0.28)
    title_text(s, "Core Capabilities — Access & Users")
    card_grid(
        s,
        [
            (
                "Authentication & Access",
                [
                    "CIMS login with PKCE for all guards",
                    "Role mapping and server-side sessions",
                    "Auto-provision trainees only",
                    "Pending approval & profile completion",
                ],
            ),
            (
                "Dashboards",
                [
                    "Role-specific dashboards",
                    "Admin stats and performance links",
                    "Trainer course/trainee metrics",
                    "Trainee courses, hours, badges",
                ],
            ),
            (
                "User & Master Data",
                [
                    "Search/filter/activate trainees",
                    "Approve registrations & assign ranks",
                    "Manage admins, trainers & education",
                    "District/tehsil and rank-course maps",
                ],
            ),
        ],
    )
    slides.append(s)

    s = blank(prs)
    fill_bg(s, prs)
    corner_accents(s, prs)
    brand_header(s)
    dots(s, 6.0, 0.28)
    title_text(s, "Learning Delivery & Assessment")
    card_grid(
        s,
        [
            (
                "Course Management",
                [
                    "Bilingual courses, categories & topics",
                    "Modules, lessons, quizzes & duration",
                    "Trainer & rank assignment",
                    "Enrollment due dates & expiry rules",
                ],
            ),
            (
                "Content & Media",
                [
                    "Videos, PDFs, docs & subtitles",
                    "Media library with S3/MinIO",
                    "Protected streaming access",
                    "Block unauthorized media navigation",
                ],
            ),
            (
                "Quizzes & Progress",
                [
                    "MCQ/fill-in quizzes with CSV upload",
                    "Time limits, passing scores, auto-grade",
                    "Lesson/course progress tracking",
                    "Completion when criteria are met",
                ],
            ),
        ],
    )
    slides.append(s)

    s = blank(prs)
    fill_bg(s, prs)
    corner_accents(s, prs)
    brand_header(s)
    dots(s, 6.0, 0.28)
    title_text(s, "Certification, Analytics & Engagement")
    card_grid(
        s,
        [
            (
                "Certificates & Badges",
                [
                    "Auto-issue on eligibility",
                    "Public verify by certificate hash",
                    "Template & signature settings",
                    "Gamification badges for milestones",
                ],
            ),
            (
                "Reports & Analytics",
                [
                    "Admin overview with charts",
                    "Trainee/course progress filters",
                    "District, tehsil, rank dimensions",
                    "CSV export for operations",
                ],
            ),
            (
                "Alerts & Support",
                [
                    "Admin alerts & deadline notices",
                    "Notification history & read status",
                    "Support tickets, chat & forums",
                    "System and profile settings",
                ],
            ),
        ],
    )
    slides.append(s)

    slides.append(
        workflow_slide(
            prs,
            "Application Workflows — Access & Learning",
            [
                (
                    "Login",
                    [
                        "Select login mode and redirect to CIMS (PKCE)",
                        "Exchange code, resolve user/role, store session",
                        "Route to dashboard, profile, or pending approval",
                    ],
                ),
                (
                    "Trainee Onboarding",
                    [
                        "Authenticate via CIMS and provision local user",
                        "Complete profile fields and documents",
                        "Admin approves; trainee gains course access",
                    ],
                ),
                (
                    "Course Creation",
                    [
                        "Admin enters course structure and bilingual info",
                        "Attach media, modules, lessons and quizzes",
                        "Validate, save, assign and activate course",
                    ],
                ),
                (
                    "Trainee Learning",
                    [
                        "Enroll/continue an active non-expired course",
                        "Complete lessons; system tracks progress & time",
                        "Pass quizzes; certificate issued on eligibility",
                    ],
                ),
            ],
        )
    )

    slides.append(
        workflow_slide(
            prs,
            "Application Workflows — Assessment & Operations",
            [
                (
                    "Course Expiry",
                    [
                        "Due date = enrollment date + duration (days)",
                        "Overdue incomplete courses marked expired",
                        "Removed from active lists; content access blocked",
                    ],
                ),
                (
                    "Quiz Attempt",
                    [
                        "Start attempt with selected questions",
                        "Finalize on submit or timeout; auto-grade",
                        "Show results; allow retry where enabled",
                    ],
                ),
                (
                    "Certificate",
                    [
                        "Check lessons complete and quizzes passed",
                        "Issue certificate with hash; mark course done",
                        "Trainee views certificate; public verify by hash",
                    ],
                ),
                (
                    "Reports & Notifications",
                    [
                        "Admin filters reports and exports CSV",
                        "Alerts and deadline notifications generated",
                        "Users view and mark notifications as read",
                    ],
                ),
            ],
        )
    )

    slides.append(slide_support(prs))
    slides.append(slide_tech(prs))
    slides.append(slide_closing(prs))

    # Apply transitions last so they sit correctly after cSld/clrMapOvr
    kinds = [
        ("fade", "l"),
        ("push", "u"),
        ("wipe", "l"),
        ("cover", "r"),
        ("fade", "l"),
        ("push", "l"),
        ("wipe", "u"),
        ("cover", "d"),
        ("fade", "l"),
        ("push", "d"),
        ("wipe", "r"),
        ("cover", "l"),
        ("fade", "l"),
        ("push", "u"),
        ("wipe", "l"),
        ("cover", "r"),
        ("fade", "l"),
        ("push", "l"),
    ]
    for i, slide in enumerate(prs.slides):
        kind, direction = kinds[i % len(kinds)]
        set_transition(slide, kind, direction)

    out = Path("/workspace/PLRA_LMS_Presentation.pptx")
    # write to temp then replace to avoid partial corrupt writes
    tmp = Path("/workspace/PLRA_LMS_Presentation.tmp.pptx")
    prs.save(str(tmp))
    count = validate_pptx(str(tmp))
    tmp.replace(out)
    print(f"Saved {out} with {count} slides (validated)")
    return str(out)


if __name__ == "__main__":
    build()
